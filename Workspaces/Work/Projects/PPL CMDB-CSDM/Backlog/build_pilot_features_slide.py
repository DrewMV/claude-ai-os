"""
Build the "Implementation Structure — Pilot Features" slide.

Recommends 1640805 (D10) and 1640819 (D11) as the pilot pair for the three-stage
feature structure, and shows why the sequencing between them matters.

Type and palette match PI4-Objectives-Features-2026-10-14.pptx.
Source: feature-implementation-structure.md + pi4-implementation-plan-appendix.md
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = "Implementation-Structure-Pilot-Features-2026-10-15.pptx"

NAVY = RGBColor(0x1E, 0x27, 0x61)
LIFE = RGBColor(0x6D, 0x2E, 0x46)
GOLD = RGBColor(0x9C, 0x6F, 0x15)
INK = RGBColor(0x1A, 0x1A, 0x2E)
GREY = RGBColor(0x6B, 0x72, 0x80)
RULE = RGBColor(0xD5, 0xD9, 0xE2)
SHADE = RGBColor(0xED, 0xEF, 0xF4)
RED = RGBColor(0xC0, 0x39, 0x2B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def tint(c, pct):
    return RGBColor(*(int(v + (255 - v) * (1 - pct)) for v in (c[0], c[1], c[2])))


COL_W = 1.37
COL = {k: 3.20 + i * COL_W for i, k in enumerate(
    ["4.1", "4.2", "4.3", "4.4", "4.5", "4.6", "4.7"])}
SPRINTS = [("4.1", "Oct 28"), ("4.2", "Nov 11"), ("4.3", "Nov 25"),
           ("4.4", "Dec 9"), ("4.5", "Dec 23"), ("4.6", "Jan 6"), ("4.7 IP", "Jan 20")]

GATE_X = 12.58

# (deliverable, name, ado, req, build, vr, margin, gate_colour)
ROWS = [
    ("D10", "Server Build Lifecycle Process Flow", "1640805",
     (COL["4.2"], COL_W), (COL["4.3"], COL_W * 2), (COL["4.6"], COL_W),
     "14 days", LIFE),
    ("D11", "Server Decommission Lifecycle Flow", "1640819",
     (COL["4.3"], COL_W), (COL["4.4"], COL_W * 3), (COL["4.7"], 1.00),
     "0 days", RED),
]
ROW_T = [2.38, 2.96]

CARDS = [
    (0.50, "Owned by you", LIFE,
     "Both features are M. Vazquez's. No new process imposed on A. Griffis, who is "
     "at 50% allocation and already carries 6 of 12 features."),
    (4.72, "Stage 3 gets a real test", LIFE,
     "Lifecycle process flows change how server builds and decommissions actually "
     "happen — genuine OCM, training and Class Manager sign-off, not a thin POC release."),
    (8.94, "Two runs, not two pilots", LIFE,
     "D10's spike closes Nov 24; D11's opens Nov 25. Fix what is awkward after the "
     "first run and apply it to the second, inside the same PI."),
]


def textbox(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for text, size, bold, colour in line:
            r = p.add_run()
            r.text = text
            r.font.name = "Calibri"
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = colour
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75, rounded=False):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        s.adjustments[0] = 0.08
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    s.text_frame.word_wrap = True
    return s


prs = Presentation()
prs.slide_width = Inches(13.3333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

tb = slide.shapes.add_textbox(Inches(0.50), Inches(0.33), Inches(12.30), Inches(0.55))
r = tb.text_frame.paragraphs[0].add_run()
r.text = "Implementation Structure — Pilot Features"
r.font.name, r.font.size, r.font.bold = "Cambria", Pt(31), True
r.font.color.rgb = NAVY

textbox(slide, 0.52, 0.90, 12.30, 0.30, [[(
    "Recommend the server lifecycle pair  ·  run the structure once, adjust, run it "
    "again  ·  both gate Feb 2, the lowest-value invoice month in the PI",
    11.5, False, GREY)]])

textbox(slide, 0.50, 6.95, 12.30, 0.30, [[(
    "PPL CMDB-CSDM  |  Feature Implementation Structure — Pilot  |  Oct 15, 2026",
    10, False, GREY)]])

# sprint grid
rect(slide, COL["4.5"], 1.55, COL_W, 1.78, fill=SHADE)
textbox(slide, 0.50, 1.60, 2.60, 0.50, [[("Pilot feature", 10, True, GREY)]])
for key, start in SPRINTS:
    x = COL[key.split()[0]]
    rect(slide, x, 1.55, 0.008, 1.78, fill=RULE)
    textbox(slide, x + 0.10, 1.60, COL_W, 0.50,
            [[(key, 10, True, NAVY)], [(start, 9, False, NAVY)]])
rect(slide, 3.20, 2.18, 9.59, 0.01, fill=RULE)

for (deliv, name, ado, req, build, vr, margin, gate_c), y in zip(ROWS, ROW_T):
    textbox(slide, 0.50, y, 2.60, 0.46, [
        [(deliv + "  ", 10, True, LIFE), (ado, 9, False, GREY)],
        [(name, 9.5, False, INK)],
    ])
    by = y + 0.06
    rect(slide, req[0], by, req[1], 0.26, fill=tint(LIFE, 0.30))
    rect(slide, build[0], by, build[1], 0.26, fill=LIFE)
    rect(slide, vr[0], by, vr[1], 0.26, fill=WHITE, line=LIFE, lw=1.25)
    d = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(GATE_X), Inches(by + 0.04),
                               Inches(0.17), Inches(0.17))
    d.fill.solid()
    d.fill.fore_color.rgb = gate_c
    d.line.fill.background()
    d.shadow.inherit = False

# handoff marker between D10 spike close and D11 spike open
rect(slide, COL["4.3"] - 0.006, 2.30, 0.014, 1.06, fill=GOLD)
textbox(slide, 4.30, 3.44, 3.30, 0.26,
        [[("handoff — D10 spike closes, D11 spike opens", 9.5, True, GOLD)]],
        align=PP_ALIGN.CENTER)

# phase legend
lx = 0.50
for label, kind, w in [("Spike", "tint", 0.50), ("Build & test", "solid", 1.00),
                       ("Validate & release", "outline", 1.25)]:
    if kind == "tint":
        rect(slide, lx, 3.50, 0.22, 0.11, fill=tint(LIFE, 0.30))
    elif kind == "solid":
        rect(slide, lx, 3.50, 0.22, 0.11, fill=LIFE)
    else:
        rect(slide, lx, 3.47, 0.22, 0.16, fill=WHITE, line=LIFE, lw=1.25)
    textbox(slide, lx + 0.28, 3.44, w, 0.24, [[(label, 9, False, GREY)]])
    lx += 0.28 + w + 0.16

# rationale cards
for x, title, colour, body in CARDS:
    c = rect(slide, x, 4.05, 3.87, 1.30, fill=tint(colour, 0.07),
             line=tint(colour, 0.30), rounded=True)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.14)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = title
    r.font.name, r.font.size, r.font.bold = "Calibri", Pt(11.5), True
    r.font.color.rgb = colour
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.LEFT
    r2 = p2.add_run()
    r2.text = body
    r2.font.name, r2.font.size = "Calibri", Pt(10)
    r2.font.color.rgb = INK

# watch band
w = rect(slide, 0.50, 5.60, 12.30, 0.82, fill=tint(GOLD, 0.10),
         line=tint(GOLD, 0.30), rounded=True)
tf = w.text_frame
tf.margin_left = tf.margin_right = Inches(0.20)
tf.margin_top = Inches(0.12)
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.LEFT
r = p.add_run()
r.text = "Watch"
r.font.name, r.font.size, r.font.bold = "Calibri", Pt(11.5), True
r.font.color.rgb = GOLD
p2 = tf.add_paragraph()
p2.alignment = PP_ALIGN.LEFT
for text, bold in [
    ("D11 finishes on its gate date with no validation margin, and its "
     "validate-and-release window falls in the IP iteration where capacity goes to "
     "I&A and PI-5 Planning — so the clean read on Stage 3 will come from D10's "
     "January window. ", False),
    ("Confirm Ray Reuter's availability for D10 requirements sign-off at the end of "
     "Sprint 4.2 before Nov 11.", True),
]:
    r2 = p2.add_run()
    r2.text = text
    r2.font.name, r2.font.size, r2.font.bold = "Calibri", Pt(10), bold
    r2.font.color.rgb = INK

prs.save(OUT)
print("wrote", OUT)
