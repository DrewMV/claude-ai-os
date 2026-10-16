"""
Build the feature story templates deck — 3 slides.

  1. Generic structure — the three buckets
  2. Template A — network inventory integration
  3. Template B — lifecycle process flow

Type and palette match PI4-Objectives-Features-2026-10-14.pptx.
Source: feature-story-templates.md
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = "Feature-Story-Templates-2026-10-15.pptx"

NAVY = RGBColor(0x1E, 0x27, 0x61)
TEAL = RGBColor(0x0B, 0x6E, 0x7F)
GOLD = RGBColor(0x9C, 0x6F, 0x15)
INK = RGBColor(0x1A, 0x1A, 0x2E)
GREY = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

BUCKETS = [(0.50, NAVY), (4.80, TEAL), (9.10, GOLD)]
COL_W = 3.70


def tint(c, pct):
    return RGBColor(*(int(v + (255 - v) * (1 - pct)) for v in (c[0], c[1], c[2])))


def textbox(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for text, size, bold, colour in line:
            r = p.add_run()
            r.text = text
            r.font.name = "Calibri"
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = colour
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75, dash=False):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                               Inches(w), Inches(h))
    s.adjustments[0] = 0.08
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
        if dash:
            s.line.dash_style = 4
    s.shadow.inherit = False
    s.text_frame.word_wrap = True
    return s


def chrome(prs, title, subtitle, footer):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(0.50), Inches(0.33), Inches(12.30), Inches(0.55))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = title
    r.font.name, r.font.size, r.font.bold = "Cambria", Pt(31), True
    r.font.color.rgb = NAVY
    textbox(slide, 0.52, 0.90, 12.30, 0.30, [[(subtitle, 11.5, False, GREY)]])
    textbox(slide, 0.50, 6.95, 12.30, 0.30, [[(footer, 10, False, GREY)]])
    return slide


def header(slide, x, colour, title, sub):
    h = rect(slide, x, 1.50, COL_W, 0.70, fill=colour)
    tf = h.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.name, r.font.size, r.font.bold = "Calibri", Pt(14.5), True
    r.font.color.rgb = WHITE
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = sub
    r2.font.name, r2.font.size = "Calibri", Pt(9.5)
    r2.font.color.rgb = tint(colour, 0.12)


def chips(slide, x, colour, items, top=2.40):
    for i, item in enumerate(items):
        y = top + i * 0.56
        c = rect(slide, x, y, COL_W, 0.44, fill=tint(colour, 0.10),
                 line=tint(colour, 0.32))
        tf = c.text_frame
        tf.margin_left = Inches(0.16)
        tf.margin_right = Inches(0.10)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        if isinstance(item, tuple):
            r = p.add_run()
            r.text = item[0] + "   "
            r.font.name, r.font.size, r.font.bold = "Calibri", Pt(10), True
            r.font.color.rgb = colour
            r2 = p.add_run()
            r2.text = item[1]
            r2.font.name, r2.font.size = "Calibri", Pt(10)
            r2.font.color.rgb = INK
        else:
            r = p.add_run()
            r.text = item
            r.font.name, r.font.size = "Calibri", Pt(10)
            r.font.color.rgb = INK


def band(slide, y, colour, label, runs, h=0.72):
    b = rect(slide, 0.50, y, 12.30, h, fill=tint(colour, 0.09),
             line=tint(colour, 0.28))
    tf = b.text_frame
    tf.margin_left = tf.margin_right = Inches(0.20)
    tf.margin_top = Inches(0.11)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = label
    r.font.name, r.font.size, r.font.bold = "Calibri", Pt(11.5), True
    r.font.color.rgb = colour
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.LEFT
    for text, bold in runs:
        r2 = p2.add_run()
        r2.text = text
        r2.font.name, r2.font.size, r2.font.bold = "Calibri", Pt(10), bold
        r2.font.color.rgb = INK


prs = Presentation()
prs.slide_width = Inches(13.3333)
prs.slide_height = Inches(7.5)

# ----------------------------------------------------- slide 1 — generic ---
s = chrome(prs, "Feature Structure",
           "Every feature decomposes into the same three buckets  ·  use it to create "
           "the spikes and stories under a feature in ADO",
           "PPL CMDB-CSDM  |  Feature Structure  |  Oct 15, 2026")

GENERIC = [
    ("Spike / requirements", "what must be answered before anyone builds", [
        "Access, data and dependencies understood",
        "Technical approach agreed with the team",
        "Build stories listed and sized",
        "Configuration vs. customization determined",
    ], "Exit — build stories pass DoR gates 1–4"),
    ("Build & test", "the stories that do the work", [
        "One story per increment of work",
        "Each story fits one iteration",
        "Testing inside the story",
        "Staging deploy + PO demo",
    ], "Exit — Definition of Done met"),
    ("Promote to prod", "getting it released and accepted", [
        "OCM comms and training where needed",
        "Change request and PROD deploy",
        "Evidence packaged",
        "Acceptance received",
    ], "Exit — gate evidenced"),
]

for (x, colour), (title, sub, items, exit_txt) in zip(BUCKETS, GENERIC):
    header(s, x, colour, title, sub)
    chips(s, x, colour, items)
    g = rect(s, x, 4.78, COL_W, 0.54, fill=None, line=colour, dash=True)
    tf = g.text_frame
    tf.margin_left = Inches(0.16)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = exit_txt
    r.font.name, r.font.size, r.font.bold = "Calibri", Pt(10.5), True
    r.font.color.rgb = colour

band(s, 5.60, NAVY, "How to use it", [
    ("Hold this next to a feature in ADO and create the items listed. Create every "
     "item with the “+” link from the Feature — never standalone. ", False),
    ("The bucket weights differ by feature type — see the two templates that follow.",
     True),
])

# ------------------------------------------------- slides 2 & 3 — templates ---
TEMPLATES = [
    ("Template A — Network Inventory Integration",
     "Applies to  1639206 SolarWinds  ·  1639225 Cisco Catalyst  ·  1639239 Qualys  "
     "·  Tripwire (no ADO item yet)",
     "PPL CMDB-CSDM  |  Template A — Network Inventory Integration  |  Oct 15, 2026",
     [("Spike / requirements", "3 items  ·  fat", [
         ("A-S1", "Source access and authentication"),
         ("A-S2", "Data profile and field mapping"),
         ("A-S3", "Reconciliation and precedence"),
     ]),
      ("Build & test", "5 items", [
          ("A-B1", "Configure data source connection"),
          ("A-B2", "Field mapping and transforms"),
          ("A-B3", "Identification and reconciliation rules"),
          ("A-B4", "Scheduled sync and error handling"),
          ("A-B5", "Validate against a sample set"),
      ]),
      ("Promote to prod", "1 item  ·  thin", [
          ("A-P1", "CR, PROD deploy, evidence, acceptance"),
      ])],
     GOLD, "Long pole — A-S1",
     [("Credentials and account permissions are the recurring blocker on this project. "
       "An integration POC changes no user-facing behaviour, so there is nothing to "
       "communicate and nobody to train — ", False),
      ("one promote item is enough.", True)]),

    ("Template B — Lifecycle Process Flow",
     "Applies to  1640805 Server Build Lifecycle  ·  1640819 Server Decommission "
     "Lifecycle",
     "PPL CMDB-CSDM  |  Template B — Lifecycle Process Flow  |  Oct 15, 2026",
     [("Spike / requirements", "3 items", [
         ("B-S1", "Current-state process map"),
         ("B-S2", "Target-state flow and CI state model"),
         ("B-S3", "Class Manager confirmation"),
     ]),
      ("Build & test", "5 items", [
          ("B-B1", "Configure lifecycle states and flow"),
          ("B-B2", "Triggers and state automation"),
          ("B-B3", "Notifications and approvals"),
          ("B-B4", "User entry point (catalog item)"),
          ("B-B5", "Test with representative records"),
      ]),
      ("Promote to prod", "3 items  ·  fat", [
          ("B-P1", "OCM comms and training materials"),
          ("B-P2", "Change request and PROD deploy"),
          ("B-P3", "Evidence packaged, acceptance received"),
      ])],
     GOLD, "Long pole — B-S3",
     [("Class Manager sign-off is an external dependency, not team capacity. A lifecycle "
       "flow changes how people actually work, so ", False),
      ("communication and training are delivery work, not an afterthought.", True)]),
]

for title, sub, footer, cols, bandc, bandlabel, bandruns in TEMPLATES:
    s = chrome(prs, title, sub, footer)
    for (x, colour), (ctitle, csub, items) in zip(BUCKETS, cols):
        header(s, x, colour, ctitle, csub)
        chips(s, x, colour, items)
    band(s, 5.60, bandc, bandlabel, bandruns)

prs.save(OUT)
print("wrote", OUT)
