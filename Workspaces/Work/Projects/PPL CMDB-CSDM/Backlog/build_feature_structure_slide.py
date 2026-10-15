"""
Build the "Feature Implementation Structure" appendix slide.

One standalone slide for introducing the three-stage feature skeleton. Type and
palette match PI4-Objectives-Features-2026-10-14.pptx so it sits with that deck.

Source: feature-implementation-structure.md
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = "Feature-Implementation-Structure-2026-10-14.pptx"

NAVY = RGBColor(0x1E, 0x27, 0x61)
TEAL = RGBColor(0x0B, 0x6E, 0x7F)
GOLD = RGBColor(0x9C, 0x6F, 0x15)
INK = RGBColor(0x1A, 0x1A, 0x2E)
GREY = RGBColor(0x6B, 0x72, 0x80)
RULE = RGBColor(0xD5, 0xD9, 0xE2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def tint(c, pct):
    return RGBColor(*(int(v + (255 - v) * (1 - pct)) for v in (c[0], c[1], c[2])))


COLS = [
    (0.50, NAVY, "Spike", "one per feature  ·  ADO type: Spike", [
        "Requirements gathered and confirmed in writing",
        "Technical approach agreed with the dev team",
        "Build stories listed and sized",
        "Configuration vs. customization determined",
    ], "Exit — build stories pass DoR gates 1–4"),
    (4.80, TEAL, "Build & test", "n stories  ·  ADO type: Story", [
        "One story per increment of work",
        "Each story fits one iteration",
        "Testing inside the story — no QA handoff",
        "Staging deploy + PO demo",
    ], "Exit — Definition of Done met"),
    (9.10, GOLD, "Implement", "one per feature  ·  ADO type: Story", [
        "Change request raised and scheduled",
        "Deployed to PROD",
        "OCM communications / training",
        "Evidence packaged, acceptance received",
    ], "Exit — contract / governance gate evidenced"),
]

COL_W = 3.70
CHIP_T = [2.58, 3.12, 3.66, 4.20]


def textbox(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if anchor:
        tf.vertical_anchor = anchor
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for text, size, bold, colour in line:
            r = p.add_run()
            r.text = text
            r.font.name = "Calibri"
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = colour
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75, dash=False):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                               Inches(w), Inches(h))
    s.adjustments[0] = 0.08
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
        if dash:
            s.line.dash_style = 4  # MSO_LINE_DASH_STYLE.DASH
    s.shadow.inherit = False
    s.text_frame.word_wrap = True
    return s


prs = Presentation()
prs.slide_width = Inches(13.3333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

tb = slide.shapes.add_textbox(Inches(0.50), Inches(0.33), Inches(12.30), Inches(0.55))
r = tb.text_frame.paragraphs[0].add_run()
r.text = "Feature Implementation Structure"
r.font.name, r.font.size, r.font.bold = "Cambria", Pt(31), True
r.font.color.rgb = NAVY

textbox(slide, 0.52, 0.90, 12.30, 0.30, [[(
    "Every feature decomposes the same way  ·  three stages, three dated exit gates  "
    "·  proposed, for agreement at PI Planning", 11.5, False, GREY)]])

textbox(slide, 0.50, 6.95, 12.30, 0.30, [[(
    "PPL CMDB-CSDM  |  Feature Implementation Structure  |  Proposed Oct 14, 2026",
    10, False, GREY)]])

for x, colour, title, sub, chips, exit_txt in COLS:
    hdr = rect(slide, x, 1.50, COL_W, 0.78, fill=colour)
    tf = hdr.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.name, r.font.size, r.font.bold = "Calibri", Pt(15), True
    r.font.color.rgb = WHITE
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = sub
    r2.font.name, r2.font.size = "Calibri", Pt(9.5)
    r2.font.color.rgb = tint(colour, 0.12)

    for chip, cy in zip(chips, CHIP_T):
        c = rect(slide, x, cy, COL_W, 0.44, fill=tint(colour, 0.10),
                 line=tint(colour, 0.35))
        tf = c.text_frame
        tf.margin_left = Inches(0.16)
        tf.margin_right = Inches(0.10)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        r = tf.paragraphs[0].add_run()
        r.text = chip
        r.font.name, r.font.size = "Calibri", Pt(10.5)
        r.font.color.rgb = INK

    g = rect(slide, x, 4.88, COL_W, 0.56, fill=None, line=colour, dash=True)
    tf = g.text_frame
    tf.margin_left = Inches(0.16)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = tf.paragraphs[0].add_run()
    r.text = exit_txt
    r.font.name, r.font.size, r.font.bold = "Calibri", Pt(10.5), True
    r.font.color.rgb = colour

for ax in (4.26, 8.56):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(ax), Inches(1.77),
                               Inches(0.28), Inches(0.24))
    a.fill.solid()
    a.fill.fore_color.rgb = RULE
    a.line.fill.background()
    a.shadow.inherit = False

# the Definition of Done boundary
b1 = rect(slide, 0.50, 5.70, 8.00, 0.38, fill=RGBColor(0xF2, 0xF4, 0xF7))
b1.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
b1.text_frame.margin_left = Inches(0.16)
r = b1.text_frame.paragraphs[0].add_run()
r.text = "Covered by the current Definition of Done"
r.font.name, r.font.size, r.font.bold = "Calibri", Pt(10.5), True
r.font.color.rgb = GREY

b2 = rect(slide, 9.10, 5.70, 3.70, 0.38, fill=tint(GOLD, 0.14))
b2.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
b2.text_frame.margin_left = Inches(0.16)
r = b2.text_frame.paragraphs[0].add_run()
r.text = "Not covered today"
r.font.name, r.font.size, r.font.bold = "Calibri", Pt(10.5), True
r.font.color.rgb = GOLD

textbox(slide, 0.50, 6.26, 12.30, 0.50, [[
    ("The Definition of Done ends at staging plus PO demo. Contract acceptance needs "
     "PROD plus customer sign-off — so a feature can be Done and still not be "
     "delivered. ", 10.5, False, GREY),
    ("The Implement story is what closes that gap, and carries its own acceptance "
     "criteria.", 10.5, True, GREY),
]])

prs.save(OUT)
print("wrote", OUT)
