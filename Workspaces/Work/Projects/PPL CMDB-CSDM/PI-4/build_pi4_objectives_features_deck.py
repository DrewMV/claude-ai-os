"""Build PI4-Objectives-Features-2026-10-16.pptx from the 2026-10-15 deck.

Applies the ADO reconciliation of 2026-10-16 (PO direction, Manuel Vazquez):

  1639206 SolarWinds : PI-3 (3.6) -> PI-4 Sprint 4.1, work ends Nov 10, 2026.

The 2026-10-15 deck treated SolarWinds as a PI-3 deliverable and carried an open
item saying ADO's PI-4 dates were wrong. That is reversed: ADO is correct that
SolarWinds is PI-4 / 4.1. Its ADO *target date* of Nov 24, 2026 is wrong, however
— the correct target is Nov 10, 2026 (confirmed by the PO). The deck therefore
shows Nov 10 and flags the ADO fix as an open item.

  D1 / Dec 31, 2026 gate is unchanged -> margin Nov 10 to Dec 31 = 51 days.

Every other feature in the deck already matched ADO and is left untouched.
PI-4 feature count goes 13 -> 14; Sprint 4.1 load 2 -> 3; program board 13 -> 14
rows (row pitch tightened from 0.310 to 0.286in to keep the board clear of the
legend at 6.36in).

The 2026-10-15 deck is left untouched. The deck itself lives in ~/Documents,
outside the vault, alongside its predecessors. This script edits a *copy* of SRC
in place rather than rebuilding, so every unchanged shape keeps its formatting.

    python build_pi4_objectives_features_deck_1016.py
    python validate_pi4_objectives_features_deck.py

Requires python-pptx.
"""
import copy
import shutil
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

SRC = Path(r"C:\Users\manuel.b.vazquez\Documents\PI4-Objectives-Features-2026-10-15.pptx")
DST = Path(r"C:\Users\manuel.b.vazquez\Documents\PI4-Objectives-Features-2026-10-16.pptx")

OLD_DATE, NEW_DATE = "Oct 15, 2026", "Oct 16, 2026"

INK = RGBColor(0x1A, 0x1A, 0x2E)
NAVY = RGBColor(0x1E, 0x27, 0x61)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

shutil.copyfile(SRC, DST)
prs = Presentation(DST)
S = list(prs.slides)

log = []


def note(msg):
    log.append(msg)


# ---------------------------------------------------------------- helpers
def shapes_by_name(slide):
    return {sh.name: sh for sh in slide.shapes}


def find_tb(slide, needle):
    """First textbox whose full text contains needle."""
    for sh in slide.shapes:
        if sh.has_text_frame and needle in sh.text_frame.text:
            return sh
    raise KeyError(f"textbox containing {needle!r} not found")


def replace_in_runs(shape, old, new):
    """Replace old->new inside individual runs (keeps all formatting)."""
    hits = 0
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if old in r.text:
                r.text = r.text.replace(old, new)
                hits += 1
    return hits


def set_para(shape, idx, text):
    """Set paragraph idx's text, keeping the first run's formatting."""
    p = shape.text_frame.paragraphs[idx]
    assert p.runs, f"paragraph {idx} has no runs (wanted to write {text!r})"
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""


def set_cell(cell, text):
    """Set a single-run cell's text, preserving its run formatting."""
    runs = [r for p in cell.text_frame.paragraphs for r in p.runs]
    assert runs, f"no runs in cell (wanted to write {text!r})"
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def next_id(slide):
    ids = [0]
    for el in slide.shapes._spTree.iter():
        if el.tag.endswith("}cNvPr"):
            try:
                ids.append(int(el.get("id")))
            except (TypeError, ValueError):
                pass
    return max(ids) + 1


def clone(slide, shape, new_name):
    el = copy.deepcopy(shape._element)
    slide.shapes._spTree.append(el)
    new = slide.shapes[-1]
    nv = el.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}cNvPr")
    if nv is None:  # pic/sp variants
        for sub in el.iter():
            if sub.tag.endswith("}cNvPr"):
                nv = sub
                break
    nv.set("id", str(next_id(slide)))
    nv.set("name", new_name)
    return new


# ================================================================ SLIDE 1
s1 = S[0]

# stat tiles: PI-4 feature count, and the integrations split
for sh in s1.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() == "13":
        set_para(sh, 0, "14")
        break
else:
    raise KeyError("slide 1 stat tile '13' not found")

set_para(find_tb(s1, "Integrations"), 0, "Integrations  ·  all in PI-4")
note("S1  stat tiles: 13 -> 14 PI-4 features; '1 in PI-3' -> 'all in PI-4'")

src_note = find_tb(s1, "Source: Azure DevOps")
set_para(src_note, 0,
         "Source: Azure DevOps (A-INFOPS\\FY26\\PI4), as of Oct 16, 2026, plus Delivery Lead "
         "direction of Oct 14: evaluate all four network gear sources for relevance and "
         "performance, with the decision to proceed to implementation taken individually for each.")
set_para(src_note, 1,
         "AP00288 deliverable numbering shown here is PROVISIONAL — drawn from an earlier draft "
         "and being verified against the current SOW. SolarWinds (1639206) is confirmed as PI-4 "
         "Sprint 4.1 work ending Nov 10, 2026; all four network gear integrations now sit in PI-4.")
note("S1  source note: SolarWinds confirmed PI-4 / 4.1")

# ================================================================ SLIDE 2
s2 = S[1]
hdr2 = find_tb(s2, "ADO features")
replace_in_runs(hdr2, "13 ADO features", "14 ADO features")

badge1 = find_tb(s2, "4 integrations")
replace_in_runs(badge1, "4 integrations — 1 in PI-3, 3 in PI-4", "4 integrations — all in PI-4")
note("S2  header 13 -> 14 ADO features; objective 1 badge -> all four integrations in PI-4")

# ================================================================ FOOTERS
for sl in S[1:]:
    for sh in sl.shapes:
        if sh.has_text_frame and OLD_DATE in sh.text_frame.text and "PPL CMDB-CSDM  |" in sh.text_frame.text:
            replace_in_runs(sh, OLD_DATE, NEW_DATE)
note(f"S2-S6 footer date -> {NEW_DATE}")

# ================================================================ SLIDE 3
s3 = S[2]
t3 = next(sh.table for sh in s3.shapes if sh.has_table)

# R2 SolarWinds — promote from greyed "delivers in PI-3" context row to a live PI-4 row
set_cell(t3.cell(2, 1), "Network Inventory Integration — SolarWinds")
set_cell(t3.cell(2, 2), "4.1")
set_cell(t3.cell(2, 3), "Nov 10, 2026")

for ci in range(6):
    cell = t3.cell(2, ci)
    cell.fill.solid()
    cell.fill.fore_color.rgb = WHITE
    for p in cell.text_frame.paragraphs:
        for r in p.runs:
            r.font.color.rgb = NAVY if ci == 4 else INK
            r.font.italic = False
note("S3  1639206 SolarWinds: PI-3 · 3.6 -> 4.1, ends Nov 10, 2026 (row un-greyed, de-italicised)")

set_para(find_tb(s3, "All four sources evaluated"), 0,
         "All four sources evaluated for relevance and performance; decision to implement taken "
         "individually — all four now scheduled in PI-4.   "
         "*Deliverable numbering provisional, pending the current AP00288.")
note("S3  subtitle: all four integrations now in PI-4")

# ================================================================ SLIDE 4
s4 = S[3]
t4 = next(sh.table for sh in s4.shapes if sh.has_table)


def set_sprint_cell(cell, features):
    """features = [(ado_id, bold_name, tail)] ; cell has one paragraph per feature,
    each with exactly three runs."""
    tf = cell.text_frame
    txBody = tf._txBody
    template = copy.deepcopy(tf.paragraphs[0]._p)
    while len(tf.paragraphs) > len(features):
        txBody.remove(tf.paragraphs[-1]._p)
    while len(tf.paragraphs) < len(features):
        txBody.append(copy.deepcopy(template))
    for p, (ado, name, tail) in zip(tf.paragraphs, features):
        runs = p.runs
        assert len(runs) == 3, f"expected 3 runs, got {len(runs)}"
        runs[0].text = f"{ado}  —  "
        runs[1].text = name
        runs[2].text = tail


F_SOLAR = ("1639206", "Network Inventory Integration — SolarWinds",
           "   ·  4.1  ·  ends Nov 10, 2026  ·  D1 gate Dec 31, 2026")
F_MAPS1 = ("1646517", "15 Priority Service Maps — Built & Validated (1 OF 3)",
           "   ·  4.1  ·  ends Nov 10, 2026  ·  D3 gate Nov 30, 2026")
F_CORR = ("1639335", "CMDB Correctness (Staleness, Orphans, Duplication)",
          "   ·  4.1–4.2  ·  ends Nov 24, 2026  ·  D8 gate Dec 31, 2026")

# only Sprint 4.1 changes; 4.2-4.7 already match ADO
set_cell(t4.cell(1, 2), "3")
set_sprint_cell(t4.cell(1, 3), [F_SOLAR, F_MAPS1, F_CORR])
note("S4  sprint load 4.1: 2 -> 3 (SolarWinds added); 4.2=4  4.3=3  4.4=1  4.5=0  4.6=3  4.7 IP=0  (14 features)")

foot4 = find_tb(s4, "Shows PI-4 work only")
runs4 = [r for p in foot4.text_frame.paragraphs for r in p.runs]
runs4[0].text = (
    "All four network gear integrations now sit in PI-4 — SolarWinds moved from PI-3 into 4.1, ending "
    "Nov 10, 2026. Cisco Catalyst moved from 4.3–4.4 to 4.2 and Qualys from 4.6–4.7 to 4.3; Tripwire now "
    "has an ADO item (1653575) in 4.6. Adoption Campaign (D7) and Server Decommission Lifecycle (D11) are "
    "the only work carrying through the holiday sprint, and no PI-4 feature extends into 4.7 IP."
)
for r in runs4[1:]:
    r.text = ""
note("S4  footnote rewritten — SolarWinds now in-PI, no PI-3 carve-out")

# ================================================================ SLIDE 5
s5 = S[4]
t5 = next(sh.table for sh in s5.shapes if sh.has_table)

# D1 SolarWinds: work now completes in PI-4, Nov 10 -> Dec 31 gate = 51 days
set_cell(t5.cell(1, 5), "Nov 10, 2026")
set_cell(t5.cell(1, 6), "✓ 51 days")
for r in [r for p in t5.cell(1, 5).text_frame.paragraphs for r in p.runs]:
    r.font.color.rgb = INK  # was muted grey while the row read "PI-3 · 3.6"
note("S5  D1 (1639206): PI-3 · 3.6 -> Nov 10, 2026; margin 'ahead of PI-4' -> 51 days")

open_items = find_tb(s5, "Open items")
oi_runs = [r for p in open_items.text_frame.paragraphs for r in p.runs]
oi_runs[0].text = "Open items — resolve before PI Planning, Nov 4, 2026"
oi_runs[1].text = (
    "1.  Direction received Oct 14: evaluate all four network gear sources for relevance and performance, "
    "with the decision to proceed taken individually for each. The new SOW is expected to match; numbering "
    "below is from an earlier draft and is still to be verified.    2.  SolarWinds (D1) and Cisco Catalyst "
    "(D2) are contracted with fixed gates, so their decision to proceed is effectively made by the contract "
    "— only Qualys (1639239) and Tripwire (1653575) carry a genuine go/no-go, and neither has a deliverable "
    "number.    3.  ADO 1639206 (SolarWinds) shows a target date of Nov 24, 2026; the correct target is "
    "Nov 10, 2026 (Sprint 4.1) as shown here — ADO needs updating to match."
)
for r in oi_runs[2:]:
    r.text = ""
note("S5  open item 3 inverted — ADO is right about PI-4/4.1, wrong about the Nov 24 target date")

# ================================================================ SLIDE 6
s6 = S[5]
set_para(find_tb(s6, "diamond marks the AP00288 gate"), 0,
         "PI-4 work only  ·  diamond marks the AP00288 gate  ·  all 14 features shown, "
         "SolarWinds (D1) starting in 4.1")
note("S6  subtitle: PI-3 exclusion removed, 14 features")

TEAL, MAROON, GOLD = "0B6E7F", "6D2E46", "9C6F15"
NAVY_HEX = "1E2761"

# label, ado, colour, bar left/width, gate x, lead left/width
ROWS = [
    ("D1",  "SolarWinds",                       "1639206", NAVY_HEX, 4.75, 1.15,  9.94,  5.90, 4.11),
    ("D3",  "15 Maps (1 of 3)",                 "1646517", TEAL,     4.75, 1.15,  7.39,  5.90, 1.56),
    ("D8",  "CMDB Correctness",                 "1639335", MAROON,   4.75, 2.30,  9.94,  7.05, 2.96),
    ("D4",  "15 Maps (2 of 3)",                 "1642237", TEAL,     5.90, 1.15,  9.94,  7.05, 2.96),
    ("D9",  "Non-Discoverable CI Interface",    "1640795", MAROON,   5.90, 1.15,  9.94,  7.05, 2.96),
    ("D2",  "Cisco Catalyst",                   "1639225", NAVY_HEX, 5.90, 1.15, 12.65,  7.05, 5.67),
    ("D7",  "Adoption Campaign",                "1639308", GOLD,     5.90, 5.75, 12.40, 11.65, 0.82),
    ("D6",  "Stakeholder Notifications",        "1642240", TEAL,     7.05, 1.15,  9.94,  8.20, 1.81),
    ("TBC", "Qualys",                           "1639239", NAVY_HEX, 7.05, 1.15, 12.65,  8.20, 4.52),
    ("D10", "Server Build Lifecycle",           "1640805", MAROON,   7.05, 2.30, 12.65,  9.35, 3.37),
    ("D11", "Server Decommission Lifecycle",    "1640819", MAROON,   8.20, 3.45, 12.65, 11.65, 1.07),
    ("D5",  "15 Maps (3 of 3)",                 "1642332", TEAL,    10.50, 1.15, 12.65, 11.65, 1.07),
    ("D12", "Service Map Update Catalog Item",  "1640844", MAROON,  10.50, 1.15, 12.65, 11.65, 1.07),
    ("TBC", "Tripwire",                         "1653575", NAVY_HEX,10.50, 1.15, 12.65, 11.65, 1.07),
]

by_name = shapes_by_name(s6)

# locate each existing row label textbox by the ADO id in its text
labels = {}
for sh in s6.shapes:
    if sh.has_text_frame:
        txt = sh.text_frame.text
        for _, _, ado, *_ in ROWS:
            if ado in txt and "  " in txt:
                labels[ado] = sh

# SolarWinds row: clone the Cisco Catalyst row (same objective colour, same bar width)
labels["1639206"] = clone(s6, labels["1639225"], "Text-solarwinds")
by_name["bar-1639206"] = clone(s6, by_name["bar-1639225"], "bar-1639206")
by_name["gate-1639206"] = clone(s6, by_name["gate-1639225"], "gate-1639206")
by_name["lead-1639206"] = clone(s6, by_name["lead-1639225"], "lead-1639206")

# 14 rows have to fit the same band (2.20in to 5.92in) that held 13 — tighten the pitch
Y0, STEP = 2.20, 3.72 / (len(ROWS) - 1)
for i, (dcode, name, ado, hexcol, bl, bw, gx, ll, lw) in enumerate(ROWS):
    ly = Y0 + i * STEP
    by = ly + 0.07
    rgb = RGBColor.from_string(hexcol)

    lbl = labels[ado]
    lbl.top = Inches(ly)
    runs = [r for p in lbl.text_frame.paragraphs for r in p.runs]
    runs[0].text = f"{dcode}  "
    runs[0].font.color.rgb = rgb
    runs[1].text = name
    runs[2].text = f"   {ado}"

    bar = by_name[f"bar-{ado}"]
    bar.left, bar.top, bar.width = Inches(bl), Inches(by), Inches(bw)
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb
    bar.line.color.rgb = rgb

    gate = by_name[f"gate-{ado}"]
    gate.left, gate.top = Inches(gx), Inches(by + 0.03)
    gate.fill.solid()
    gate.fill.fore_color.rgb = rgb
    gate.line.color.rgb = rgb

    lead = by_name[f"lead-{ado}"]
    lead.left, lead.top, lead.width = Inches(ll), Inches(by + 0.10), Inches(lw)
    lead.line.color.rgb = rgb

note(f"S6  program board 13 -> 14 rows; D1 SolarWinds added in 4.1; row pitch 0.310 -> {STEP:.3f}in")

prs.save(DST)
print(f"WROTE {DST}\n")
for line in log:
    print(" •", line)
